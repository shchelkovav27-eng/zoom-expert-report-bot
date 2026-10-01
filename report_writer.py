"""
Черновик «Отчёта эксперта после встречи»: нейросеть пишет разделы по расшифровке,
затем текст вписывается в копию шаблона templates/report_template.docx (оформление сохраняется).
"""

from __future__ import annotations

import json
import re
from copy import deepcopy

import anthropic
from docx import Document
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

from config import ANTHROPIC_API_KEY, LLM_MODEL, REPORT_TEMPLATE

TEXT_COLOR = RGBColor(0x1A, 0x2B, 0x34)

# Разделы формы: ключ → (заголовок в шаблоне, что в нём писать — из подсказок формы)
SECTIONS = {
    "overview": ("Как прошла встреча",
                 "Короткий обзор: с чего начали, чем занимались, чем закончили. Как команда работала — "
                 "включённо, формально, спорила ли между собой."),
    "closed": ("Что закрыли",
               "Какие поля Дорожной карты заполнены и в каком состоянии: черновик, обсуждено с командой "
               "или утверждено. Если что-то осталось незакрытым — прямо, что именно."),
    "decisions": ("Главные решения и формулировки",
                  "Что решила сама организация. Итоговые формулировки — по возможности дословно, в кавычках."),
    "strengths": ("Что хорошо и на что можно опереться",
                  "Сильные места: удачные формулировки, включённые люди, готовые данные, разумные решения."),
    "risks": ("Риски и тревожные сигналы",
              "Где команда себя переоценивает, где решение принято формально, где не хватает людей "
              "или данных, где есть внутреннее разногласие. Отметь, если ситуация повторяется."),
    "agreements": ("О чём договорились до следующей встречи",
                   "Домашняя работа с именами и сроками, недостающие цифры и документы, кто что находит "
                   "и к какому числу."),
    "help": ("Нужна ли помощь программы",
             "Что стоит решить не на уровне организации: вопрос к руководителю программы, тема для "
             "мастермайнда, необходимость подключить другого эксперта."),
}

SYSTEM = """Ты помогаешь эксперту программы сопровождения НКО. По расшифровке записи встречи \
составь черновик «Отчёта эксперта после встречи».

Как писать:
- Связным текстом, как эксперт рассказал бы коллеге, который будет вести эту организацию вместо него. \
Это не таблица и не протокол. Весь отчёт — примерно страница: каждый раздел 1–2 коротких абзаца.
- Только то, что есть в расшифровке и чате. Ничего не выдумывай. Если по разделу в записи ничего нет — \
так и напиши: «В записи не прозвучало».
- Итоговые формулировки решений приводи дословно, в кавычках.
- Договорённости — с именами и сроками; если имя или срок не прозвучали, отметь это.
- Риски — профессиональное суждение по тому, что видно из записи.
- Спикеры в расшифровке могут быть не подписаны или подписаны условно — определяй людей по смыслу.
- Пиши по-русски, без Markdown (без звёздочек и решёток). Абзацы разделяй пустой строкой.

Что писать в каждом поле ответа:
participants — кто участвовал и кого не было из нужных людей (одной-двумя фразами);
{keys}."""


def _schema() -> dict:
    fields = ["participants", *SECTIONS]
    return {
        "type": "object",
        "properties": {k: {"type": "string"} for k in fields},
        "required": fields,
        "additionalProperties": False,
    }


CONTINUATION = """

Эта встреча — продолжение: по тому же модулю и номеру встреча уже была {prev_date}, команда не успела \
всё отработать. Ниже дан отчёт прошлой встречи. Опирайся на него: в разделе «Что закрыли» отдельно скажи, \
что добавилось к прошлому разу, а что так и осталось незакрытым; в «Рисках» отметь, если проблема повторяется. \
Не переписывай прошлый отчёт: пиши про эту встречу, ссылаясь на прошлую там, где это важно."""


def draft_report(meta: dict, transcript: str, chat: str = "", previous_report: str = "") -> dict:
    if not ANTHROPIC_API_KEY:
        raise RuntimeError("Не задан ANTHROPIC_API_KEY в .env")
    keys = ";\n".join(f"{k} — «{title}»: {hint}" for k, (title, hint) in SECTIONS.items())
    user = (
        f"Организация: {meta['org']}\nЭксперт: {meta['expert']}\n"
        f"Модуль и встреча: {meta['block']}\nДата: {meta['date']}\nНазвание в Zoom: {meta['topic']}\n\n"
        + (f"ЧАТ ВСТРЕЧИ:\n{chat}\n\n" if chat.strip() else "")
        + f"РАСШИФРОВКА:\n{transcript}"
    )
    system = SYSTEM.format(keys=keys)
    if previous_report.strip():
        system += CONTINUATION.format(prev_date=meta.get("prev_date", "раньше"))
        user += f"\n\nОТЧЁТ ПРОШЛОЙ ВСТРЕЧИ:\n{previous_report}"
    client = anthropic.Anthropic(api_key=ANTHROPIC_API_KEY, timeout=900, max_retries=3)
    response = client.messages.create(
        model=LLM_MODEL,
        max_tokens=16000,
        system=system,
        messages=[{"role": "user", "content": user}],
        # Ответ строго по схеме — сразу готовые поля отчёта
        output_config={"format": {"type": "json_schema", "schema": _schema()}},
        # Если основная модель откажется отвечать, запрос сам уйдёт на запасную
        extra_headers={"anthropic-beta": "server-side-fallback-2026-07-01"},
        extra_body={"fallbacks": "default"},
    )
    if response.stop_reason == "refusal":
        raise RuntimeError("Claude отказался составлять отчёт по этой записи")
    if response.stop_reason == "max_tokens":
        raise RuntimeError("Отчёт не поместился в лимит ответа")
    text = next(b.text for b in response.content if b.type == "text")
    data = json.loads(text)
    return {k: (data.get(k) or "В записи не прозвучало.").strip() for k in ["participants", *SECTIONS]}


# ---------- заполнение шаблона ----------

def _clear(p: Paragraph) -> None:
    for r in list(p.runs):
        r._r.getparent().remove(r._r)


def _add_run(p: Paragraph, text: str, *, italic: bool = False, size: float = 10.5) -> None:
    run = p.add_run(text)
    run.bold, run.italic = False, italic
    run.font.size = Pt(size)
    run.font.color.rgb = TEXT_COLOR


def _write_paragraphs(p: Paragraph, text: str) -> None:
    """Заменяет подсказку в разделе на текст; абзацы — отдельными параграфами с тем же отступом."""
    parts = [x.strip() for x in re.split(r"\n\s*\n|\n", text) if x.strip()] or ["—"]
    _clear(p)
    blank = deepcopy(p._p)
    p.paragraph_format.space_after = Pt(6)
    _add_run(p, parts[0])
    anchor = p._p
    for part in parts[1:]:
        el = deepcopy(blank)
        anchor.addnext(el)
        anchor = el
        np = Paragraph(el, p._parent)
        np.paragraph_format.space_after = Pt(6)
        _add_run(np, part)


def _remove(p: Paragraph) -> None:
    p._p.getparent().remove(p._p)


def fill_template(meta: dict, data: dict, out_path: str) -> str:
    doc = Document(REPORT_TEMPLATE)
    header = {
        "Организация:": meta["org"],
        "Эксперт:": meta["expert"],
        "Блок и номер встречи, дата:": f"{meta['block']}, {meta['date']}",
        "Кто участвовал": data["participants"],
    }
    titles = {title: key for key, (title, _) in SECTIONS.items()}
    section, hint_pending = None, False

    for p in list(doc.paragraphs):
        t = p.text.strip()
        if t.startswith(("Отчёт пишется", "Пишите так")):  # инструкция для эксперта — в отчёте не нужна
            _remove(p)
            continue
        label = next((lb for lb in header if t.startswith(lb)), None)
        if label:
            _add_run(p, " " + header[label])
            continue
        if t in titles:
            section, hint_pending = titles[t], True
            continue
        if t.startswith("Отчёт заполнил"):
            _clear(p)
            _add_run(p, f"Черновик подготовлен автоматически по записи встречи {meta.get('date_short', meta['date'])}. "
                        f"{meta['expert']}, проверьте, дополните своим суждением и поставьте дату.",
                     italic=True, size=10)
            section = None
            continue
        if section and hint_pending and t:
            _write_paragraphs(p, data[section])
            hint_pending = False
        elif section and not t:
            _remove(p)  # пустые строки «для письма от руки»

    for table in list(doc.tables):  # памятка «Чтобы расшифровка была пригодной» — для встречи, не для отчёта
        if table.cell(0, 0).text.strip().startswith("ЧТОБЫ РАСШИФРОВКА"):
            table._tbl.getparent().remove(table._tbl)

    doc.save(out_path)
    return out_path
